#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sistema de Gestión Tributaria - Extractor de Precisión SRI
VERSIÓN 3.0:
  - Casilleros detectados por su RECUADRO DE COLOR (celeste / naranja / azul),
    ya no por "número seguido de número". Los códigos que son solo parte de un
    texto o fórmula (620+621, "trasládese el campo 615"...) se ignoran.
  - Valores tomados tal cual del PDF (sin redondeos tipo 0.00615 -> 0.01)
  - Reporte especial: Casilleros × Meses/Ubicación

Características:
  • Extracción de formularios SRI (101, 102, 103, 104, 107, 115, etc.)
  • Vouchear Talón Resumen ATS (Anexo Transaccional)
  • Validaciones automáticas
  • Exportación a Excel (2 modos: pestañas o columnas)
  • Interfaz gráfica drag-and-drop

Autor: Roberto Robles © Unidad Técnica PBP
"""

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, simpledialog
from tkinterdnd2 import TkinterDnD, DND_FILES
import datetime
import json
import os
import re
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ============= CONFIGURACIÓN =============
RUTA_TESSERACT = r"C:\Users\rrobles\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"
if os.path.exists(RUTA_TESSERACT):
    pytesseract.pytesseract.tesseract_cmd = RUTA_TESSERACT

LICENCIA = "eyJle-HAiOi-IyMDI-2LTA5-LTI0I-n0uwy-I2_4P-Nt5oG-k0iX3-CIPmr-p6HAX-Q0YmN-sa4t3-oZsYI-g"
DURACION_MESES = 6
RUTA_CONFIG = os.path.join(os.path.expanduser("~"), "voucheo_config.json")

# Listas globales
DATOS_PROCESADOS = []
DATOS_ATS_PROCESADOS = []

# ============= FUNCIONES AUXILIARES BÁSICAS =============
def safe_messagebox(tipo, titulo, mensaje):
    dummy = tk.Toplevel()
    dummy.withdraw()
    dummy.attributes("-topmost", True)
    if tipo == "info":
        messagebox.showinfo(titulo, mensaje, parent=dummy)
    elif tipo == "warning":
        messagebox.showwarning(titulo, mensaje, parent=dummy)
    elif tipo == "error":
        messagebox.showerror(titulo, mensaje, parent=dummy)
    dummy.destroy()

# ============= LICENCIA =============
def verificar_licencia():
    if os.path.exists(RUTA_CONFIG):
        with open(RUTA_CONFIG, "r") as f:
            config = json.load(f)
        if "fecha_activacion" in config:
            fecha_ini = datetime.datetime.strptime(config["fecha_activacion"], "%Y-%m-%d")
            dias_restantes = DURACION_MESES * 30 - (datetime.datetime.now() - fecha_ini).days
            if dias_restantes > 0:
                return dias_restantes
    licencia_ingresada = simpledialog.askstring("Licencia", "Introduce tu licencia de activación:")
    if licencia_ingresada != LICENCIA:
        safe_messagebox("error", "Licencia inválida", "Licencia incorrecta.")
        return False
    with open(RUTA_CONFIG, "w") as f:
        json.dump({"fecha_activacion": datetime.datetime.now().strftime("%Y-%m-%d")}, f)
    return DURACION_MESES * 30

def limpiar_rutas_arrastre(data_string):
    pattern = re.compile(r'\{([^}]+)\}|(\S+)')
    return [
        os.path.normpath(m.group(1) if m.group(1) else m.group(2))
        for m in pattern.finditer(data_string)
        if (m.group(1) if m.group(1) else m.group(2)).lower().endswith('.pdf')
    ]

# ============= LECTURA DE TEXTO Y TABLAS DEL PDF =============
def ocr_pdf_completo(ruta_pdf):
    texto = ""
    try:
        doc = fitz.open(ruta_pdf)
        for pagina in doc:
            pix = pagina.get_pixmap(dpi=300)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            try:
                texto += pytesseract.image_to_string(img, lang="spa") + "\n"
            except Exception:
                texto += pytesseract.image_to_string(img) + "\n"
        doc.close()
    except Exception:
        pass
    return texto

def extraer_texto_y_tablas(ruta_pdf):
    texto_completo = ""
    tablas = []
    try:
        with pdfplumber.open(ruta_pdf) as pdf:
            for pagina in pdf.pages:
                texto_pagina = pagina.extract_text()
                if texto_pagina:
                    texto_completo += texto_pagina + "\n"
                try:
                    for tabla in pagina.extract_tables():
                        tablas.append(tabla)
                except Exception:
                    pass
    except Exception as e:
        safe_messagebox("error", "Error de Extracción", f"No se pudo abrir el PDF con pdfplumber: {e}")

    if len(texto_completo.strip()) < 40:
        texto_ocr = ocr_pdf_completo(ruta_pdf)
        if texto_ocr:
            texto_completo += texto_ocr

    return texto_completo, tablas

# ============= NORMALIZACIÓN (usada por el ATS) =============
class SeparadorPrecision:
    """Normaliza números de texto plano (Talón ATS)."""

    DECIMALES_VALIDOS = (2, 4)
    UMBRAL_VALOR_PEGADO = 0.005

    @classmethod
    def normalizar(cls, valor_str):
        if not valor_str:
            return None

        v = valor_str.replace('$', '').replace(' ', '').strip()

        if v in ("", "-", "--", "N/A", "n/a"):
            return None

        if ',' in v and '.' in v:
            if v.rfind(',') > v.rfind('.'):
                entero, _, decimal = v.rpartition(',')
                if len(decimal) not in cls.DECIMALES_VALIDOS:
                    return None
                v = entero.replace('.', '') + '.' + decimal
            else:
                entero, _, decimal = v.rpartition('.')
                if len(decimal) not in cls.DECIMALES_VALIDOS:
                    return None
                v = entero.replace(',', '') + '.' + decimal
        elif ',' in v:
            entero, _, decimal = v.rpartition(',')
            if len(decimal) in cls.DECIMALES_VALIDOS:
                v = entero + '.' + decimal
            elif len(decimal) == 3 and entero.replace('.', '').isdigit():
                v = entero.replace(',', '') + '.00'
            else:
                return None
        elif '.' in v:
            entero, _, decimal = v.rpartition('.')
            if len(decimal) not in cls.DECIMALES_VALIDOS:
                return None

        try:
            valor_float = float(v)
            if 0 < abs(valor_float) < cls.UMBRAL_VALOR_PEGADO:
                return None
            if valor_float == 0:
                return "0.00"
            return f"{valor_float:.2f}"
        except (ValueError, TypeError):
            return None

def normalizar_valor_mejorado(valor_str):
    return SeparadorPrecision.normalizar(valor_str)

# ============= EXTRACCIÓN DE CASILLEROS POR RECUADRO DE COLOR =============
# En los formularios SRI un CASILLERO es el número que está dentro de un recuadro
# de color (celeste, naranja, azul oscuro...). El VALOR es lo que está a su derecha,
# en la misma fila. Cualquier otro número (620+621, "trasládese el campo 615",
# un 0.00 pegado a un código...) es solo texto y se IGNORA.
PATRON_CODIGO_CASILLERO = re.compile(r'^\d{3,4}$')
PATRON_VALOR_NUMERICO = re.compile(r'^-?\$?\d[\d,]*(?:\.\d+)?$|^-?\$?\.\d+$')
PATRON_RESPUESTA_TEXTO = re.compile(r'^(?:SI|SÍ|NO|NO APLICA)$', re.IGNORECASE)
MAX_ANCHO_CAJA = 130
MAX_ALTO_CAJA = 90

def _es_color_casillero(color):
    """Celeste, naranja, azules = casillero. Blanco y grises = celdas de valor/texto."""
    if not isinstance(color, (list, tuple)) or len(color) != 3:
        return False
    if min(color) > 0.97:
        return False
    return (max(color) - min(color)) > 0.08

def _limpiar_valor(texto):
    return texto.replace('$', '').replace(',', '').strip()

def _extraer_casilleros_por_recuadro(ruta_pdf):
    casilleros = {}
    origen = {}
    registro_debug = []
    paginas_sin_vector = []   # páginas escaneadas/imagen: se resuelven por OCR
    try:
        with pdfplumber.open(ruta_pdf) as pdf:
            for num_pag, pagina in enumerate(pdf.pages, 1):
                cajas = [r for r in pagina.rects
                         if _es_color_casillero(r.get('non_stroking_color'))
                         and r['width'] < MAX_ANCHO_CAJA and r['height'] < MAX_ALTO_CAJA]
                palabras = pagina.extract_words(x_tolerance=1.5, y_tolerance=2)
                if not cajas:
                    if len(palabras) < 20:
                        paginas_sin_vector.append(num_pag)
                    continue

                def caja_de(w):
                    cx = (w['x0'] + w['x1']) / 2
                    cy = (w['top'] + w['bottom']) / 2
                    for b in cajas:
                        if b['x0'] - 1 <= cx <= b['x1'] + 1 and b['top'] - 1 <= cy <= b['bottom'] + 1:
                            return b
                    return None

                codigos = []
                for w in palabras:
                    if PATRON_CODIGO_CASILLERO.fullmatch(w['text']):
                        b = caja_de(w)
                        if b is not None:
                            codigos.append((w, b))

                def en_fila(w, b):
                    cy = (w['top'] + w['bottom']) / 2
                    return b['top'] - 2 <= cy <= b['bottom'] + 2

                for w, b in codigos:
                    # límite derecho: el siguiente casillero de la misma fila
                    limite = min([c[1]['x0'] for c in codigos
                                  if c[1] is not b and c[1]['x0'] > b['x1'] - 1 and en_fila(c[0], b)] or [1e9])
                    derecha = sorted(
                        [x for x in palabras
                         if x['x0'] >= b['x1'] - 1 and x['x0'] < limite and en_fila(x, b)],
                        key=lambda x: x['x0'])
                    valor = ""
                    tipo = "vacio"
                    for x in derecha:
                        if PATRON_CODIGO_CASILLERO.fullmatch(x['text']):
                            continue  # otro código suelto, no es valor
                        if PATRON_VALOR_NUMERICO.fullmatch(x['text']):
                            valor = _limpiar_valor(x['text'])
                            tipo = "numero"
                            break
                    if tipo == "vacio":
                        textos = [x['text'] for x in derecha]
                        dos = " ".join(textos[:2])
                        if textos and PATRON_RESPUESTA_TEXTO.fullmatch(dos):
                            valor, tipo = dos, "texto"
                        elif textos and PATRON_RESPUESTA_TEXTO.fullmatch(textos[0]):
                            valor, tipo = textos[0], "texto"
                    codigo = w['text']
                    if codigo not in casilleros or (casilleros[codigo] == "" and valor != ""):
                        casilleros[codigo] = valor
                        origen[codigo] = "recuadro" if tipo != "vacio" else "recuadro_sin_valor"
                        registro_debug.append({'pagina': num_pag, 'codigo': codigo,
                                               'valor': valor, 'origen': origen[codigo]})
    except Exception:
        pass
    return casilleros, origen, registro_debug, paginas_sin_vector

# ============= OCR PARA PDF ESCANEADOS (imagen) =============
# Mismo criterio que el método vectorial: casillero = número de 3-4 dígitos sobre
# fondo de color; valor = número a su derecha en la misma fila.
OCR_DPI = 300
PATRON_VALOR_OCR = re.compile(r'^-?\$?\d[\d,]*\.\d{2,6}$|^-?\$?\d{1,15}$')

def _fondo_es_color(img_rgb, x0, y0, x1, y1):
    """Color de fondo dominante de la caja de la palabra (ignora el trazo del texto)."""
    w, h = img_rgb.size
    x0, y0 = max(0, int(x0) - 3), max(0, int(y0) - 1)
    x1, y1 = min(w, int(x1) + 3), min(h, int(y1) + 1)
    if x1 <= x0 or y1 <= y0:
        return False
    muestra = img_rgb.crop((x0, y0, x1, y1)).quantize(colors=4).convert("RGB")
    cuentas = muestra.getcolors(maxcolors=16) or []
    if not cuentas:
        return False
    _, color = max(cuentas, key=lambda t: t[0])
    return (max(color) - min(color)) > 0.08 * 255 and min(color) < 0.97 * 255

def _palabras_ocr(img):
    """OCR normal + OCR invertido (texto blanco sobre azul oscuro). Une sin duplicar."""
    from PIL import ImageOps
    gris = img.convert("L")
    palabras = []
    for variante in (gris, ImageOps.invert(gris)):
        try:
            d = pytesseract.image_to_data(variante, lang="spa", config="--psm 11",
                                          output_type=pytesseract.Output.DICT)
        except Exception:
            d = pytesseract.image_to_data(variante, config="--psm 11",
                                          output_type=pytesseract.Output.DICT)
        for i, t in enumerate(d["text"]):
            t = (t or "").strip()
            if not t:
                continue
            w = {"text": t, "x0": d["left"][i], "x1": d["left"][i] + d["width"][i],
                 "top": d["top"][i], "bottom": d["top"][i] + d["height"][i]}
            cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
            duplicada = any(abs(cx - (o["x0"] + o["x1"]) / 2) < 6 and abs(cy - (o["top"] + o["bottom"]) / 2) < 6
                            for o in palabras)
            if not duplicada:
                palabras.append(w)
    return palabras

PATRON_DECIMAL_OCR = re.compile(r'^-?\d[\d,]*\.\d{2,6}$')
OPERADORES_FORMULA = re.compile(r'^[\(\-\+=\*/]+\)?$|^(?:campo|campos|casillero|casilleros|y|más|mas|menos)$', re.IGNORECASE)

def _limpiar_valor_ocr(texto):
    """Quita ruido de bordes de celda (— | _ ~ ') delante del número; conserva '-' (negativo)."""
    t = texto.strip()
    t = re.sub(r'^[\u2014\u2013\u2012|_~\'`.,:;]+', '', t)
    t = re.sub(r'[|_~\'`:;]+$', '', t)
    return _limpiar_valor(t)

def _es_parte_de_formula(w, palabras, k_dpi):
    """620+621, '399 - 898', '(trasládese campo 429) 482': el código va pegado a un operador/texto."""
    alto = max(w["bottom"] - w["top"], 1)
    cy = (w["top"] + w["bottom"]) / 2
    fila = [x for x in palabras if x is not w and abs((x["top"] + x["bottom"]) / 2 - cy) <= 0.8 * alto]
    izq = [x for x in fila if x["x1"] <= w["x0"] + 2 and w["x0"] - x["x1"] <= 5 * k_dpi]
    der = [x for x in fila if x["x0"] >= w["x1"] - 2 and x["x0"] - w["x1"] <= 5 * k_dpi]
    if izq:
        t = max(izq, key=lambda x: x["x1"])["text"]
        if OPERADORES_FORMULA.fullmatch(t):
            return True
    if der:
        t = min(der, key=lambda x: x["x0"])["text"]
        if re.fullmatch(r'[\-\+=\*/]+', t):
            return True
    return False

def _banda_tiene_color(img, x0, y0, x1, y1):
    w, h = img.size
    x0, y0, x1, y1 = max(0, int(x0)), max(0, int(y0)), min(w, int(x1)), min(h, int(y1))
    if x1 - x0 < 4 or y1 - y0 < 2:
        return False
    muestra = img.crop((x0, y0, x1, y1)).quantize(colors=6).convert("RGB")
    total = (x1 - x0) * (y1 - y0)
    for cuenta, color in (muestra.getcolors(maxcolors=64) or []):
        if cuenta > 0.03 * total and (max(color) - min(color)) > 0.08 * 255 and min(color) < 0.97 * 255:
            return True
    return False

def _pixel_es_color(p):
    return (max(p) - min(p)) > 0.08 * 255 and min(p) < 0.97 * 255

def _celdas_color_en_fila(img, x0, x1, y_a, y_b, hueco_max=3):
    """Tramos horizontales de fondo coloreado (celdas) en una franja. Devuelve [(x_ini, x_fin)]."""
    w, h = img.size
    x0, x1 = max(0, int(x0)), min(w, int(x1))
    ya, yb = min(max(0, int(y_a)), h - 1), min(max(0, int(y_b)), h - 1)
    px = img.load()
    tramos, ini, ultimo, huecos = [], None, None, 0
    for x in range(x0, x1):
        col = _pixel_es_color(px[x, ya]) or _pixel_es_color(px[x, yb])
        if col:
            if ini is None:
                ini = x
            ultimo, huecos = x, 0
        elif ini is not None:
            huecos += 1
            if huecos > hueco_max:
                tramos.append((ini, ultimo))
                ini = None
    if ini is not None:
        tramos.append((ini, ultimo))
    return tramos

def _ocr_codigo_en_celda(img, x0, y0, x1, y1):
    """OCR de dígitos dentro de UNA celda coloreada (normal e invertido). Devuelve texto de 3-4 dígitos o None."""
    from PIL import ImageOps
    recorte = img.crop((int(x0) + 2, int(y0), int(x1) - 2, int(y1)))
    if recorte.width < 8 or recorte.height < 8:
        return None
    ampliado = recorte.resize((recorte.width * 2, recorte.height * 2), Image.LANCZOS).convert("L")
    for variante in (ampliado, ImageOps.invert(ampliado)):
        try:
            t = pytesseract.image_to_string(
                variante, config="--psm 7 -c tessedit_char_whitelist=0123456789").strip()
        except Exception:
            continue
        if PATRON_CODIGO_CASILLERO.fullmatch(t):
            return t
    return None

def _reocr_valor(img, x, original, margen=6):
    """
    Re-lee un monto que el OCR devolvió sin decimales (ej. 6242944 en vez de 6242944.17).
    Amplía el recorte hacia la derecha (donde estaban los decimales perdidos). Los montos del SRI
    siempre llevan 2 decimales: si el OCR pierde el punto pero entrega 2 dígitos más, se reconstruye.
    """
    from PIL import ImageOps
    base = original.lstrip("-")
    for ext in (60, 30, 0):
        recorte = img.crop((int(x["x0"]) - margen, int(x["top"]) - margen, int(x["x1"]) + ext, int(x["bottom"]) + margen))
        ampliado = recorte.resize((recorte.width * 3, recorte.height * 3), Image.LANCZOS).convert("L")
        for variante in (ampliado, ImageOps.invert(ampliado)):
            try:
                t = pytesseract.image_to_string(
                    variante, config="--psm 7 -c tessedit_char_whitelist=0123456789.,-").strip()
            except Exception:
                continue
            t = _limpiar_valor_ocr(t)
            if PATRON_DECIMAL_OCR.fullmatch(t) and t.lstrip("-").split(".")[0] == base:
                return t
            solo = t.lstrip("-").replace(".", "")
            if solo.startswith(base) and len(solo) == len(base) + 2 and solo.isdigit():
                return ("-" if original.startswith("-") else "") + base + "." + solo[-2:]
    return None

def _extraer_casilleros_ocr(ruta_pdf, paginas):
    casilleros = {}
    origen = {}
    registro_debug = []
    if not paginas:
        return casilleros, origen, registro_debug
    try:
        doc = fitz.open(ruta_pdf)
    except Exception:
        return casilleros, origen, registro_debug
    for num_pag in paginas:
        try:
            pix = doc[num_pag - 1].get_pixmap(dpi=OCR_DPI)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            palabras = _palabras_ocr(img)
        except Exception:
            continue
        k_dpi = OCR_DPI / 72.0
        codigos = [w for w in palabras
                   if PATRON_CODIGO_CASILLERO.fullmatch(w["text"])
                   and _fondo_es_color(img, w["x0"], w["top"], w["x1"], w["bottom"])
                   and not _es_parte_de_formula(w, palabras, k_dpi)]
        reclamados = set()
        for w in codigos:
            alto = max(w["bottom"] - w["top"], 1)
            cy = (w["top"] + w["bottom"]) / 2
            en_fila = lambda x: abs((x["top"] + x["bottom"]) / 2 - cy) <= 0.8 * alto
            limite = min([c["x0"] for c in codigos if c is not w and c["x0"] > w["x1"] and en_fila(c)] or [1e9])
            derecha = sorted([x for x in palabras
                              if x["x0"] >= w["x1"] and x["x0"] < limite and en_fila(x)],
                             key=lambda x: x["x0"])
            valor = ""
            dudoso = False
            for x in derecha:
                t = _limpiar_valor_ocr(x["text"].replace("O", "0").replace("o", "0"))
                if PATRON_VALOR_OCR.fullmatch(t) and not PATRON_CODIGO_CASILLERO.fullmatch(t):
                    if "." not in t and len(t.lstrip("-")) > 3:
                        t_fix = _reocr_valor(img, x, t)   # monto sin decimales: probablemente OCR truncado
                        if t_fix:
                            t = t_fix
                        else:
                            dudoso = True
                    valor = t
                    reclamados.add(id(x))
                    break
            if not valor:
                for x in derecha:
                    t = _limpiar_valor_ocr(x["text"])
                    if t == "0":
                        valor = "0"
                        reclamados.add(id(x))
                        break
            codigo = w["text"]
            if codigo not in casilleros or (casilleros[codigo] == "" and valor != ""):
                casilleros[codigo] = valor
                origen[codigo] = ("ocr_revisar" if dudoso else "ocr_recuadro") if valor else "ocr_sin_valor"
                registro_debug.append({"pagina": num_pag, "codigo": codigo,
                                       "valor": valor, "origen": origen[codigo]})
        # --- Recuperación: valores sin casillero (el OCR no leyó el código, ej. fondo oscuro)
        for v in palabras:
            if id(v) in reclamados:
                continue
            tv = _limpiar_valor_ocr(v["text"])
            if not PATRON_DECIMAL_OCR.fullmatch(tv):
                continue
            alto = max(v["bottom"] - v["top"], 1)
            cy = (v["top"] + v["bottom"]) / 2
            x_ini = max(0, v["x0"] - 140 * k_dpi)
            y0, y1 = cy - 0.9 * alto, cy + 0.9 * alto
            if not _banda_tiene_color(img, x_ini, cy - 0.3 * alto, v["x0"] - 2, cy + 0.3 * alto):
                continue
            conocidos = [c for c in codigos if abs((c["top"] + c["bottom"]) / 2 - cy) <= 0.8 * alto]
            tramos = _celdas_color_en_fila(img, x_ini, v["x0"] - 2, cy - 0.65 * alto, cy + 0.65 * alto)
            codigo = None
            for t0, t1 in sorted(tramos, key=lambda t: -t[1]):  # de derecha a izquierda
                ancho = t1 - t0
                if ancho < 6 * k_dpi or ancho > 130 * k_dpi:
                    continue
                if any(c["x0"] - 8 <= (t0 + t1) / 2 <= c["x1"] + 8 or t0 <= c["x0"] <= t1 for c in conocidos):
                    break  # el tramo más cercano ya es un casillero conocido: el valor es de ese
                cand = _ocr_codigo_en_celda(img, t0, cy - 0.9 * alto, t1, cy + 0.9 * alto)
                if cand:
                    codigo = cand
                break
            if codigo and codigo not in casilleros:
                casilleros[codigo] = tv
                origen[codigo] = "ocr_recuperado"
                reclamados.add(id(v))
                registro_debug.append({"pagina": num_pag, "codigo": codigo,
                                       "valor": tv, "origen": "ocr_recuperado"})
    doc.close()
    return casilleros, origen, registro_debug

def extraer_casilleros_mejorado(ruta_pdf, texto_completo, tablas, valores_excluir):
    """
    1. Recuadros de color vectoriales (PDF original del SRI, método principal)
    2. OCR con detección de fondo de color (páginas escaneadas / imagen)
    3. Texto y tablas: último recurso si no se detectó ningún recuadro
    """
    casilleros, origen, debug, paginas_imagen = _extraer_casilleros_por_recuadro(ruta_pdf)
    if paginas_imagen:
        c_ocr, o_ocr, d_ocr = _extraer_casilleros_ocr(ruta_pdf, paginas_imagen)
        for cod, val in c_ocr.items():
            if cod not in casilleros:
                casilleros[cod] = val
                origen[cod] = o_ocr[cod]
        debug.extend(d_ocr)
    if casilleros:
        return casilleros, origen, debug

    patron_par = re.compile(r'(?<![\d.,+])(\d{3,4})\s+(-?\d+\.\d{2})(?![\d])')
    for match in patron_par.finditer(texto_completo):
        cod, val = match.group(1), match.group(2)
        if cod not in valores_excluir and cod not in casilleros:
            casilleros[cod] = val
            origen[cod] = "ocr_fallback"

    for tabla in tablas:
        for fila in tabla:
            if not fila or len(fila) < 2:
                continue
            celdas = [(c.strip() if c else "") for c in fila]
            cod = next((c for c in celdas[:2] if PATRON_CODIGO_CASILLERO.fullmatch(c)), None)
            if not cod or cod in casilleros or cod in valores_excluir:
                continue
            for c in reversed(celdas):
                c_limpio = _limpiar_valor(c)
                if re.fullmatch(r'-?\d+\.\d{2}', c_limpio):
                    casilleros[cod] = c_limpio
                    origen[cod] = "tabla"
                    break
    return casilleros, origen, debug

# ============= EXTRACCIÓN DE METADATOS =============
def extraer_metadatos(texto_completo, lineas):
    meta = {
        "Formulario": "No detectado",
        "Numero_Formulario": "No detectado",
        "RUC": "No detectado",
        "Razon_Social": "No detectado",
        "Periodo": "No detectado",
        "Numero_Serial": "No detectado",
        "Codigo_Verificador": "No detectado",
        "Fecha_Presentacion": "No detectado",
    }
    meses = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
             "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

    for i, linea in enumerate(lineas):
        linea_clean = linea.strip()
        if not linea_clean:
            continue
        mayus = linea_clean.upper()

        if meta["Formulario"] == "No detectado" and (
            "DECLARACIÓN" in mayus or "DECLARACION" in mayus or "OBLIGACIÓN TRIBUTARIA" in mayus
            or "IMPUESTO" in mayus
        ):
            meta["Formulario"] = linea_clean

        match_num_form = re.search(r'FORMULARIO\s*N?[o°º]?\.?\s*(\d{2,4})', mayus)
        if match_num_form and meta["Numero_Formulario"] == "No detectado":
            meta["Numero_Formulario"] = match_num_form.group(1)

        match_ruc = re.search(r'\b\d{13}\b', linea_clean)
        if match_ruc and meta["RUC"] == "No detectado":
            meta["RUC"] = match_ruc.group(0)

        match_serial = re.search(r'\b\d{10,15}\b', linea_clean)
        if match_serial and meta["Numero_Serial"] == "No detectado":
            contexto_previo = lineas[max(0, i - 1)].upper()
            if "SERIAL" in mayus or "SERIAL" in contexto_previo or i > len(lineas) - 15:
                meta["Numero_Serial"] = match_serial.group(0)

        match_verificador = re.search(r'\b([A-Z]{2,10}\d{6,})\b', mayus)
        if match_verificador and meta["Codigo_Verificador"] == "No detectado":
            meta["Codigo_Verificador"] = match_verificador.group(1)

        if meta["Periodo"] == "No detectado" and any(mes in mayus for mes in meses):
            if "PERÍODO" in mayus or "PERIODO" in mayus or "EJERCICIO" in mayus or len(linea_clean.split()) <= 5:
                meta["Periodo"] = linea_clean

        if meta["Periodo"] == "No detectado":
            match_periodo_num = re.search(r'\bPER[IÍ]ODO[^\d]{0,15}(\d{1,2})\s*[\/\-]\s*(\d{4})\b', mayus)
            if match_periodo_num:
                meta["Periodo"] = f"{match_periodo_num.group(1)}/{match_periodo_num.group(2)}"

        match_fecha = re.search(r'\b(\d{1,2}[\/\-]\d{1,2}[\/\-]\d{4})\b', linea_clean)
        if match_fecha and meta["Fecha_Presentacion"] == "No detectado" and (
            "FECHA" in mayus or "PRESENT" in mayus or "RECEPCI" in mayus
        ):
            meta["Fecha_Presentacion"] = match_fecha.group(1)

        if ("RAZÓN SOCIAL" in mayus or "RAZON SOCIAL" in mayus) and meta["Razon_Social"] == "No detectado":
            despues_dos_puntos = re.split(r'RAZ[OÓ]N SOCIAL\s*:?\s*', linea_clean, flags=re.IGNORECASE)
            candidato = despues_dos_puntos[-1].strip() if len(despues_dos_puntos) > 1 else ""
            if not candidato and i + 1 < len(lineas):
                candidato = lineas[i + 1].strip()
            meta["Razon_Social"] = candidato.replace('"', '').replace(',', '').strip() or "No detectado"

    return meta

def extraer_mes_inteligente(texto_completo, periodo_str):
    """
    Devuelve el período como MM/AAAA (mensual) o 'Anual AAAA' (formulario anual, ej. 101).
    Solo mira el texto del PERÍODO (no el resto del formulario).
    """
    meses_nombres = {
        "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4,
        "MAYO": 5, "JUNIO": 6, "JULIO": 7, "AGOSTO": 8,
        "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12
    }
    candidatos = []
    if periodo_str and periodo_str != "No detectado":
        candidatos.append(periodo_str)
    if texto_completo:
        for m in re.finditer(r'PER[IÍ]ODO\s+FISCAL\s*:?\s*([^\n]+)|(?:^|\n)\s*(?:MES|PER[IÍ]ODO)\s*:?\s*([^\n]+)',
                             texto_completo, re.IGNORECASE):
            candidatos.append(m.group(1) or m.group(2))
        for m in re.finditer(r'(?:^|\n)\s*A[ÑN]O\s*:?\s*(20\d{2})', texto_completo, re.IGNORECASE):
            candidatos.append("AÑO " + m.group(1))

    for cand in candidatos:
        up = cand.upper()
        m_num = re.search(r'\b(\d{1,2})\s*[\/\-]\s*(20\d{2})\b', up)
        if m_num and 1 <= int(m_num.group(1)) <= 12:
            return f"{int(m_num.group(1)):02d}/{m_num.group(2)}"
        for nombre, num in meses_nombres.items():
            if re.search(r'\b' + nombre + r'\b', up):
                m_anio = re.search(r'\b(20\d{2})\b', up)
                if m_anio:
                    return f"{num:02d}/{m_anio.group(1)}"
        m_anual = re.search(r'\bA[ÑN]O\s*:?\s*(20\d{2})\b', up)
        if m_anual:
            return f"Anual {m_anual.group(1)}"
    return "Desconocido"

# ============= VALIDACIÓN =============
def _es_numero(valor):
    try:
        float(valor)
        return True
    except (TypeError, ValueError):
        return False

def validar_datos_mejorado(meta, casilleros):
    val = {
        "ruc_ok": bool(re.fullmatch(r'\d{13}', meta.get("RUC", ""))),
        "serial_ok": meta.get("Numero_Serial", "No detectado") != "No detectado",
        "razon_ok": meta.get("Razon_Social", "No detectado") != "No detectado",
        "periodo_ok": meta.get("Periodo", "No detectado") != "No detectado",
        "formulario_ok": meta.get("Formulario", "No detectado") != "No detectado",
        "casilleros_count": len(casilleros),
    }
    val["casilleros_ok"] = val["casilleros_count"] > 0
    # casilleros detectados (recuadro) pero sin valor en el PDF
    val["casilleros_sin_valor"] = len([v for v in casilleros.values() if v == ""])
    # se mantiene la clave por compatibilidad con la interfaz/reportes
    val["valores_sospechosos"] = 0
    val["codigos_duplicados"] = False
    val["completo"] = val["ruc_ok"] and val["serial_ok"] and val["casilleros_ok"]
    return val

# ============= FUNCIÓN PRINCIPAL DE EXTRACCIÓN =============
def extraer_datos_formulario_sri(ruta_pdf):
    texto_completo, tablas = extraer_texto_y_tablas(ruta_pdf)
    lineas = texto_completo.split("\n")
    meta = extraer_metadatos(texto_completo, lineas)

    valores_excluir = set()
    if meta["RUC"] != "No detectado":
        valores_excluir.add(meta["RUC"])
    if meta["Numero_Serial"] != "No detectado":
        valores_excluir.add(meta["Numero_Serial"])

    casilleros, origen, debug = extraer_casilleros_mejorado(
        ruta_pdf, texto_completo, tablas, valores_excluir
    )

    validacion = validar_datos_mejorado(meta, casilleros)
    validacion["casilleros_a_revisar"] = sum(1 for o in origen.values() if o in ("ocr_revisar", "ocr_sin_valor"))
    mes_formulario = extraer_mes_inteligente(texto_completo, meta.get("Periodo", ""))

    return {
        "archivo": ruta_pdf,
        "Metadatos": meta,
        "Casilleros": casilleros,
        "Origen": origen,
        "Validacion": validacion,
        "Debug": debug,
        "Mes": mes_formulario,
    }

# ============= DEBUGGING =============
def generar_reporte_debug(datos):
    reporte = []
    reporte.append(f"\n{'='*70}")
    reporte.append(f"REPORTE DE EXTRACCIÓN: {datos['Metadatos'].get('Numero_Formulario', 'N/A')}")
    reporte.append(f"{'='*70}")

    reporte.append(f"\n📊 RESUMEN:")
    reporte.append(f"  • RUC: {datos['Metadatos'].get('RUC', 'N/A')}")
    reporte.append(f"  • Periodo: {datos['Metadatos'].get('Periodo', 'N/A')}")
    reporte.append(f"  • Casilleros encontrados: {datos['Validacion']['casilleros_count']}")
    reporte.append(f"  • Casilleros sin valor en el PDF: {datos['Validacion'].get('casilleros_sin_valor', 0)}")
    reporte.append(f"  • Validación: {'✔ COMPLETO' if datos['Validacion']['completo'] else '⚠ INCOMPLETO'}")

    reporte.append(f"\n📋 CASILLEROS POR ORIGEN:")
    origenes = {}
    for cod, ori in datos['Origen'].items():
        origenes[ori] = origenes.get(ori, 0) + 1
    for ori, count in sorted(origenes.items()):
        reporte.append(f"  • {ori}: {count} casilleros")

    reporte.append(f"\n📝 DETALLE DE CASILLEROS:")
    for codigo in sorted(datos['Casilleros'].keys(), key=int):
        valor = datos['Casilleros'][codigo]
        origen = datos['Origen'].get(codigo, 'desconocido')
        reporte.append(f"  {codigo:>6} = {valor:>14} ({origen})")

    return "\n".join(reporte)

# ============= PORTAPAPELES =============
def _orden_casillero(codigo):
    try:
        return (0, int(codigo))
    except (TypeError, ValueError):
        return (1, 0)

def construir_texto_portapapeles(datos):
    matriz = "CONCEPTO / CASILLERO\tVALOR\n"
    for campo, contenido in datos["Metadatos"].items():
        matriz += f"{campo}\t{contenido}\n"
    matriz += "\n"
    for codigo in sorted(datos["Casilleros"].keys(), key=_orden_casillero):
        matriz += f"Casillero {codigo}\t{datos['Casilleros'][codigo]}\n"
    return matriz

def copiar_al_portapapeles_formateado():
    datos = obtener_datos_pestana_actual()
    if not datos:
        safe_messagebox("warning", "Sin datos", "Cargue primero uno o varios PDF y seleccione una pestaña.")
        return
    matriz_excel = construir_texto_portapapeles(datos)
    ventana.clipboard_clear()
    ventana.clipboard_append(matriz_excel)
    ventana.update()
    safe_messagebox(
        "info", "Copiado Exitoso",
        f"Datos de '{os.path.basename(datos['archivo'])}' copiados al portapapeles.\n\n"
        f"MÉTRICAS DE INTEGRIDAD:\n"
        f"-----------------------------------------\n"
        f"• Casilleros detectados: {datos['Validacion']['casilleros_count']}\n"
        f"• RUC: {datos['Metadatos']['RUC']}\n"
        f"• Serial SRI: {datos['Metadatos']['Numero_Serial']}\n"
        f"-----------------------------------------\n\n"
        f"Ahora puede hacer clic en cualquier celda de Excel y presionar Ctrl+V."
    )

# ============= EXPORTACIÓN A EXCEL =============
def generar_nombre_hoja(datos, usados):
    meta = datos["Metadatos"]
    base = meta.get("Numero_Formulario", "Form")
    if base == "No detectado":
        base = "Form"
    base = re.sub(r'[\[\]\:\*\?\/\\]', '', str(base))[:10]
    periodo = meta.get("Periodo", "")
    periodo = re.sub(r'[\[\]\:\*\?\/\\]', '', str(periodo))[:12]
    ruc_corto = meta.get("RUC", "")[-4:] if meta.get("RUC", "No detectado") != "No detectado" else ""
    nombre = f"F{base}_{periodo}_{ruc_corto}".strip('_')
    nombre = re.sub(r'\s+', '_', nombre)
    if not nombre:
        nombre = "Formulario"
    nombre = nombre[:28]
    candidato = nombre
    i = 1
    while candidato in usados:
        i += 1
        candidato = f"{nombre[:25]}_{i}"
    usados.add(candidato)
    return candidato[:31]

NOMBRE_HOJA_CONSOLIDADA = "Consolidado_SRI"

def _columna_letra(indice_col):
    return openpyxl.utils.get_column_letter(indice_col)

def _exportar_modo_pestanas(wb):
    usados = set(wb.sheetnames)
    hojas_creadas = 0
    for datos in DATOS_PROCESADOS:
        nombre_hoja = generar_nombre_hoja(datos, usados)
        ws = wb.create_sheet(title=nombre_hoja)
        ws.append(["CONCEPTO", "VALOR"])
        for campo, contenido in datos["Metadatos"].items():
            ws.append([campo, contenido])
        ws.append([])
        ws.append(["CASILLERO", "VALOR"])
        for codigo in sorted(datos["Casilleros"].keys(), key=_orden_casillero):
            ws.append([codigo, datos["Casilleros"][codigo]])
        ws.column_dimensions['A'].width = 32
        ws.column_dimensions['B'].width = 22
        hojas_creadas += 1
    return f"Se agregaron {hojas_creadas} pestaña(s) (una por formulario)."

def _exportar_modo_columnas(wb):
    """Coloca cada formulario en columnas dentro de UNA sola pestaña."""
    if NOMBRE_HOJA_CONSOLIDADA in wb.sheetnames:
        ws = wb[NOMBRE_HOJA_CONSOLIDADA]
        if ws.cell(row=1, column=1).value is None:
            col_inicio = 1
        else:
            col_inicio = ws.max_column + 2
    else:
        ws = wb.create_sheet(title=NOMBRE_HOJA_CONSOLIDADA)
        col_inicio = 1

    formularios_agregados = 0
    for datos in DATOS_PROCESADOS:
        col_concepto = col_inicio
        col_valor = col_inicio + 1

        fila = 1
        ws.cell(row=fila, column=col_concepto, value="CONCEPTO")
        ws.cell(row=fila, column=col_valor, value="VALOR")
        fila += 1

        for campo, contenido in datos["Metadatos"].items():
            ws.cell(row=fila, column=col_concepto, value=campo)
            ws.cell(row=fila, column=col_valor, value=contenido)
            fila += 1

        fila += 1

        ws.cell(row=fila, column=col_concepto, value="CASILLERO")
        ws.cell(row=fila, column=col_valor, value="VALOR")
        fila += 1

        for codigo in sorted(datos["Casilleros"].keys(), key=_orden_casillero):
            ws.cell(row=fila, column=col_concepto, value=codigo)
            ws.cell(row=fila, column=col_valor, value=datos["Casilleros"][codigo])
            fila += 1

        ws.column_dimensions[_columna_letra(col_concepto)].width = 32
        ws.column_dimensions[_columna_letra(col_valor)].width = 22

        col_inicio = col_valor + 2
        formularios_agregados += 1

    return (f"Se agregaron {formularios_agregados} formulario(s) en columnas "
            f"dentro de la pestaña '{NOMBRE_HOJA_CONSOLIDADA}'.")

def exportar_a_excel():
    if not DATOS_PROCESADOS:
        safe_messagebox("warning", "Sin datos", "Cargue al menos un PDF antes de exportar.")
        return
    modo = var_modo_exportacion.get()
    usar_existente = messagebox.askyesno(
        "Exportar a Excel",
        "¿Desea agregar los datos a un archivo Excel YA EXISTENTE?\n\n"
        "Sí = seleccionar un Excel existente y agregarle los datos.\n"
        "No = crear un archivo Excel nuevo."
    )
    ruta_existente = None
    if usar_existente:
        ruta_existente = filedialog.askopenfilename(
            title="Seleccionar Excel existente",
            filetypes=[("Libros de Excel", "*.xlsx")]
        )
        if not ruta_existente:
            return
        try:
            wb = openpyxl.load_workbook(ruta_existente)
        except Exception as e:
            safe_messagebox("error", "Error", f"No se pudo abrir el Excel seleccionado: {e}")
            return
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    if modo == "pestanas":
        mensaje_resultado = _exportar_modo_pestanas(wb)
    else:
        mensaje_resultado = _exportar_modo_columnas(wb)

    ruta_guardar = filedialog.asksaveasfilename(
        title="Guardar archivo Excel",
        defaultextension=".xlsx",
        initialfile=os.path.basename(ruta_existente) if ruta_existente else "Formularios_SRI.xlsx",
        filetypes=[("Libro de Excel", "*.xlsx")]
    )
    if not ruta_guardar:
        return
    try:
        wb.save(ruta_guardar)
    except Exception as e:
        safe_messagebox("error", "Error", f"No se pudo guardar el archivo: {e}")
        return
    safe_messagebox(
        "info", "Exportación Exitosa",
        f"{mensaje_resultado}\n\nArchivo:\n{ruta_guardar}"
    )

# ============= REPORTE ESPECIAL CASILLEROS × MESES =============
def generar_reporte_especial_casilleros():
    """Casilleros en filas, meses en columnas."""
    if not DATOS_PROCESADOS:
        safe_messagebox("warning", "Sin datos", "Cargue al menos un PDF antes de generar el reporte especial.")
        return

    reporte_datos = {}
    todos_meses = set()

    for datos in DATOS_PROCESADOS:
        mes = datos.get("Mes", "Desconocido")
        todos_meses.add(mes)
        for codigo, valor in datos["Casilleros"].items():
            reporte_datos.setdefault(codigo, {})[mes] = valor

    casilleros_ordenados = sorted(reporte_datos.keys(), key=_orden_casillero)
    meses_ordenados = sorted(todos_meses)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(title="Reporte_Especial")

    ws.cell(row=1, column=1, value="Casillero / Mes")
    for col_idx, mes in enumerate(meses_ordenados, start=2):
        ws.cell(row=1, column=col_idx, value=mes)

    for row_idx, casillero in enumerate(casilleros_ordenados, start=2):
        ws.cell(row=row_idx, column=1, value=casillero)
        for col_idx, mes in enumerate(meses_ordenados, start=2):
            valor = reporte_datos[casillero].get(mes, "")
            if valor == "" or valor is None:
                ws.cell(row=row_idx, column=col_idx, value="")
            elif _es_numero(valor):
                celda = ws.cell(row=row_idx, column=col_idx, value=float(valor))
                celda.number_format = '#,##0.00'
            else:
                ws.cell(row=row_idx, column=col_idx, value=valor)

    for col_idx in range(1, len(meses_ordenados) + 2):
        ws.cell(row=1, column=col_idx).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=col_idx).fill = PatternFill(start_color="0D47A1", end_color="0D47A1", fill_type="solid")

    ws.column_dimensions['A'].width = 15
    for col_idx in range(2, len(meses_ordenados) + 2):
        ws.column_dimensions[_columna_letra(col_idx)].width = 18

    ruta_guardar = filedialog.asksaveasfilename(
        title="Guardar Reporte Especial",
        defaultextension=".xlsx",
        initialfile="Reporte_Especial_Casilleros_Meses.xlsx",
        filetypes=[("Libro de Excel", "*.xlsx")]
    )
    if not ruta_guardar:
        return

    try:
        wb.save(ruta_guardar)
        safe_messagebox(
            "info", "Reporte Especial Generado",
            f"Reporte generado exitosamente.\n\n"
            f"Casilleros: {len(casilleros_ordenados)}\n"
            f"Meses: {len(meses_ordenados)}\n\n"
            f"Archivo:\n{ruta_guardar}"
        )
    except Exception as e:
        safe_messagebox("error", "Error", f"No se pudo guardar el reporte: {e}")

# ============= ATS: TALÓN RESUMEN =============
COLOR_TITULO_ATS = "0D47A1"
COLOR_ENCABEZADO_TRANSACCION = "1565C0"
COLOR_ENCABEZADO_RETENCION_RENTA = "EF6C00"
COLOR_ENCABEZADO_RETENCION_IVA = "6A1B9A"
COLOR_FILA_PAR_ATS = "F2F6FC"
COLOR_TOTAL_ATS = "C8E6C9"
BORDE_FINO_ATS = Border(
    left=Side(style="thin", color="B0BEC5"), right=Side(style="thin", color="B0BEC5"),
    top=Side(style="thin", color="B0BEC5"), bottom=Side(style="thin", color="B0BEC5")
)

ANCLAS_ATS_ORDEN = [
    "COMPRAS",
    "VENTAS",
    "EXPORTACIONES",
    "RETENCION EN LA FUENTE DE IMPUESTO A LA RENTA",
    "RETENCION EN LA FUENTE DE IVA",
    "DECLARO QUE",
]

PATRON_FILA_COMPLETA_ATS = re.compile(
    r'^\s*(?P<codigo>\d{2,4}[A-Za-z]?)\s+(?P<descripcion>.+?)\s+(?P<registros>\d+)\s+'
    r'(?P<vals>(?:-?\d+[.,]\d{2}\s*)+)$'
)
PATRON_FILA_SOLO_CODIGO_ATS = re.compile(
    r'^\s*(?P<codigo>\d{2,4}[A-Za-z]?)\s+(?P<registros>\d+)\s+(?P<vals>(?:-?\d+[.,]\d{2}\s*)+)$'
)
PATRON_LINEA_ES_DATO_ATS = re.compile(r'^\s*\d{2,4}[A-Za-z]?\s')
PATRON_LINEA_TERMINA_DECIMAL_ATS = re.compile(r'-?\d+[.,]\d{2}\s*$')
PATRON_FILA_IVA_ATS = re.compile(
    r'^[ \t]*(?P<operacion>COMPRA|VENTA)[ \t]+(?P<concepto>Retenci[oó]n[ \t]+IVA[ \t]+(?:\d+[ \t]*%|NC))[ \t]+'
    r'(?P<valor>-?\d+[.,]\d{2})[ \t]*$',
    re.MULTILINE
)

def _limpiar_texto_multilinea_ats(txt):
    return re.sub(r'\s+', ' ', txt).strip()

def extraer_metadatos_ats(texto_completo, lineas):
    meta = {
        "Razon_Social": "No detectado",
        "RUC": "No detectado",
        "Periodo": "No detectado",
        "Fecha_Generacion": "No detectado",
        "Estado": "No detectado",
        "Secuencial_Anexo": "No detectado",
    }
    m_ruc = re.search(r'RUC:\s*(\d{10,13})', texto_completo)
    if m_ruc:
        meta["RUC"] = m_ruc.group(1)
        idx_linea_ruc = None
        for i, linea in enumerate(lineas):
            if "RUC:" in linea.upper():
                idx_linea_ruc = i
                break
        if idx_linea_ruc is not None:
            encabezados_conocidos = ("TALÓN RESUMEN", "TALON RESUMEN", "SERVICIO DE RENTAS INTERNAS", "ANEXO TRANSACCIONAL")
            for j in range(idx_linea_ruc - 1, -1, -1):
                candidato = lineas[j].strip()
                if candidato and candidato.upper() not in encabezados_conocidos:
                    meta["Razon_Social"] = candidato
                    break
    m_periodo = re.search(r'Periodo:\s*([^\n]+)', texto_completo)
    if m_periodo:
        meta["Periodo"] = m_periodo.group(1).strip()
    m_fecha = re.search(r'Fecha de Generaci[oó]n:\s*([^\n]+)', texto_completo)
    if m_fecha:
        meta["Fecha_Generacion"] = m_fecha.group(1).strip()
    m_estado = re.search(r'Estado:\s*([^\n]+)', texto_completo)
    if m_estado:
        meta["Estado"] = m_estado.group(1).strip()
    m_secuencial = re.search(r'Secuencial Anexo:\s*([^\n]+)', texto_completo)
    if m_secuencial:
        meta["Secuencial_Anexo"] = m_secuencial.group(1).strip()
    return meta

def extraer_bloque_seccion_ats(texto, nombre_seccion):
    texto_upper = texto.upper()
    nombre_upper = nombre_seccion.upper()
    m_inicio = re.search(r'\b' + re.escape(nombre_upper) + r'\b', texto_upper)
    if not m_inicio:
        return ""
    inicio = m_inicio.end()
    resto_upper = texto_upper[inicio:]
    fin_relativo = len(resto_upper)
    try:
        idx_seccion = ANCLAS_ATS_ORDEN.index(nombre_upper)
        anclas_siguientes = ANCLAS_ATS_ORDEN[idx_seccion + 1:]
    except ValueError:
        anclas_siguientes = []
    for ancla in anclas_siguientes:
        m_fin = re.search(r'\b' + re.escape(ancla) + r'\b', resto_upper)
        if m_fin and m_fin.start() < fin_relativo:
            fin_relativo = m_fin.start()
    return texto[inicio:inicio + fin_relativo]

def extraer_filas_transaccion_ats(bloque_texto):
    filas = []
    if not bloque_texto:
        return filas
    lineas = bloque_texto.split("\n")
    n = len(lineas)
    i = 0
    while i < n:
        linea = lineas[i].strip()
        if not linea:
            i += 1
            continue
        m_completo = PATRON_FILA_COMPLETA_ATS.match(linea)
        if m_completo:
            valores = [normalizar_valor_mejorado(v) for v in m_completo.group("vals").split()]
            valores = [v for v in valores if v is not None]
            if valores:
                filas.append({
                    "codigo": m_completo.group("codigo"),
                    "descripcion": _limpiar_texto_multilinea_ats(m_completo.group("descripcion")),
                    "registros": m_completo.group("registros"),
                    "valores": valores,
                })
            i += 1
            continue
        m_solo = PATRON_FILA_SOLO_CODIGO_ATS.match(linea)
        if m_solo:
            partes_desc = []
            if i > 0:
                anterior = lineas[i - 1].strip()
                if (anterior and not PATRON_LINEA_ES_DATO_ATS.match(anterior)
                        and "TOTAL:" not in anterior.upper()
                        and not PATRON_LINEA_TERMINA_DECIMAL_ATS.search(anterior)):
                    partes_desc.append(anterior)
            avanzar = 1
            if i + 1 < n:
                siguiente = lineas[i + 1].strip()
                if (siguiente and not PATRON_LINEA_ES_DATO_ATS.match(siguiente)
                        and "TOTAL:" not in siguiente.upper()
                        and not PATRON_LINEA_TERMINA_DECIMAL_ATS.search(siguiente)
                        and not PATRON_FILA_SOLO_CODIGO_ATS.match(siguiente)):
                    partes_desc.append(siguiente)
                    avanzar = 2
            valores = [normalizar_valor_mejorado(v) for v in m_solo.group("vals").split()]
            valores = [v for v in valores if v is not None]
            if valores:
                filas.append({
                    "codigo": m_solo.group("codigo"),
                    "descripcion": _limpiar_texto_multilinea_ats(" ".join(partes_desc)),
                    "registros": m_solo.group("registros"),
                    "valores": valores,
                })
            i += avanzar
            continue
        i += 1
    return filas

def extraer_filas_iva_ats(bloque_texto):
    filas = []
    if not bloque_texto:
        return filas
    for m in PATRON_FILA_IVA_ATS.finditer(bloque_texto):
        valor = normalizar_valor_mejorado(m.group("valor"))
        if valor is None:
            continue
        filas.append({
            "operacion": m.group("operacion"),
            "concepto": _limpiar_texto_multilinea_ats(m.group("concepto")),
            "valor": valor,
        })
    return filas

def extraer_total_bloque_ats(bloque_texto):
    if not bloque_texto:
        return []
    coincidencias = re.findall(r'TOTAL:\s*((?:-?\d+[.,]\d{2}\s*)+)', bloque_texto)
    if not coincidencias:
        return []
    valores = [normalizar_valor_mejorado(v) for v in coincidencias[-1].split()]
    return [v for v in valores if v is not None]

def _detectar_etiquetas_valor_ats(bloque_texto, num_columnas, etiquetas_reserva):
    etiquetas = []
    coincidencias_tarifa = re.findall(r'BI\s*tarifa\s*(?:diferente\s*0\s*%|\d+\s*%)', bloque_texto, re.IGNORECASE)
    etiquetas.extend(_limpiar_texto_multilinea_ats(t) for t in coincidencias_tarifa[:2])
    m_no_obj = re.search(r'BI\s*No\s*Objeto\s*IVA', bloque_texto, re.IGNORECASE)
    if m_no_obj:
        etiquetas.append(_limpiar_texto_multilinea_ats(m_no_obj.group(0)))
    m_valor = re.search(r'Valor\s*(?:IVA|FOB)', bloque_texto, re.IGNORECASE)
    if m_valor:
        etiquetas.append(_limpiar_texto_multilinea_ats(m_valor.group(0)))
    if len(etiquetas) >= num_columnas:
        return etiquetas[:num_columnas]
    faltan = num_columnas - len(etiquetas)
    return etiquetas + etiquetas_reserva[len(etiquetas):len(etiquetas) + faltan] + \
        [f"Valor {i+1}" for i in range(len(etiquetas) + faltan, num_columnas)]

def _procesar_seccion_transaccion_ats(texto_completo, nombre_seccion, etiquetas_reserva):
    bloque = extraer_bloque_seccion_ats(texto_completo, nombre_seccion)
    filas = extraer_filas_transaccion_ats(bloque)
    total = extraer_total_bloque_ats(bloque)
    if not filas and not total:
        return None
    num_columnas_valor = max([len(f["valores"]) for f in filas], default=len(total))
    if num_columnas_valor == 0:
        return None
    etiquetas_valor = _detectar_etiquetas_valor_ats(bloque, num_columnas_valor, etiquetas_reserva)
    return {"filas": filas, "total": total, "etiquetas_valor": etiquetas_valor}

def extraer_datos_ats(ruta_pdf):
    texto_completo, _tablas = extraer_texto_y_tablas(ruta_pdf)
    lineas = texto_completo.split("\n")
    meta = extraer_metadatos_ats(texto_completo, lineas)

    compras = _procesar_seccion_transaccion_ats(
        texto_completo, "COMPRAS",
        ["BI tarifa 0%", "BI tarifa 12%", "BI No Objeto IVA", "Valor IVA"]
    )
    ventas = _procesar_seccion_transaccion_ats(
        texto_completo, "VENTAS",
        ["BI tarifa 0%", "BI tarifa 12%", "BI No Objeto IVA", "Valor IVA"]
    )
    exportaciones = _procesar_seccion_transaccion_ats(
        texto_completo, "EXPORTACIONES",
        ["Valor FOB"]
    )
    ret_renta = _procesar_seccion_transaccion_ats(
        texto_completo, "RETENCION EN LA FUENTE DE IMPUESTO A LA RENTA",
        ["Base Imponible", "Valor Retenido"]
    )

    bloque_iva = extraer_bloque_seccion_ats(texto_completo, "RETENCION EN LA FUENTE DE IVA")
    filas_iva = extraer_filas_iva_ats(bloque_iva)
    total_iva = extraer_total_bloque_ats(bloque_iva)
    ret_iva = {"filas": filas_iva, "total": total_iva} if (filas_iva or total_iva) else None

    return {
        "archivo": ruta_pdf,
        "Metadatos": meta,
        "Compras": compras,
        "Ventas": ventas,
        "Exportaciones": exportaciones,
        "RetencionRenta": ret_renta,
        "RetencionIVA": ret_iva,
    }

# ============= ESCRITURA ATS EN EXCEL =============
def _pintar_banner_ats(ws, fila, texto, num_columnas, color_hex):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=num_columnas)
    celda = ws.cell(row=fila, column=1, value=texto)
    celda.font = Font(bold=True, color="FFFFFF", size=11)
    celda.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
    celda.alignment = Alignment(horizontal="left", vertical="center")
    for c in range(1, num_columnas + 1):
        ws.cell(row=fila, column=c).border = BORDE_FINO_ATS
    ws.row_dimensions[fila].height = 20

def _pintar_encabezado_columnas_ats(ws, fila, etiquetas, color_hex):
    for idx, etiqueta in enumerate(etiquetas, start=1):
        celda = ws.cell(row=fila, column=idx, value=etiqueta if etiqueta else None)
        celda.font = Font(bold=True, color="FFFFFF", size=9)
        celda.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE_FINO_ATS
    ws.row_dimensions[fila].height = 26

def _valor_celda_numerico_ats(valor_str):
    try:
        return float(valor_str)
    except (TypeError, ValueError):
        return valor_str if valor_str not in ("", None) else None

def _pintar_fila_datos_ats(ws, fila, valores_por_columna, es_par):
    color_fondo = COLOR_FILA_PAR_ATS if es_par else "FFFFFF"
    for idx, valor in enumerate(valores_por_columna, start=1):
        celda = ws.cell(row=fila, column=idx)
        if idx >= 4:
            celda.value = _valor_celda_numerico_ats(valor)
            if isinstance(celda.value, float):
                celda.number_format = '#,##0.00'
            celda.alignment = Alignment(horizontal="right")
        else:
            celda.value = valor if valor not in ("", None) else None
            celda.alignment = Alignment(horizontal="left", wrap_text=(idx == 2))
        celda.fill = PatternFill(start_color=color_fondo, end_color=color_fondo, fill_type="solid")
        celda.border = BORDE_FINO_ATS
        celda.font = Font(size=9)

def _pintar_fila_total_ats(ws, fila, valores_por_columna, num_columnas):
    for idx in range(1, num_columnas + 1):
        celda = ws.cell(row=fila, column=idx)
        celda.fill = PatternFill(start_color=COLOR_TOTAL_ATS, end_color=COLOR_TOTAL_ATS, fill_type="solid")
        celda.font = Font(bold=True, size=9)
        celda.border = BORDE_FINO_ATS
    for idx, valor in enumerate(valores_por_columna, start=1):
        celda = ws.cell(row=fila, column=idx)
        if idx >= 4:
            celda.value = _valor_celda_numerico_ats(valor)
            if isinstance(celda.value, float):
                celda.number_format = '#,##0.00'
            celda.alignment = Alignment(horizontal="right")
        else:
            celda.value = valor if valor not in ("", None) else None

def _escribir_seccion_ats(ws, fila_inicio, titulo, etiquetas_valor, filas_normalizadas, total_valores, color_tema, num_columnas_total=7):
    if not filas_normalizadas and not total_valores:
        return fila_inicio
    fila = fila_inicio
    _pintar_banner_ats(ws, fila, titulo, num_columnas_total, color_tema)
    fila += 1
    etiquetas_completas = ["Código", "Descripción / Concepto", "No. Registros"] + list(etiquetas_valor)
    etiquetas_completas += [""] * (num_columnas_total - len(etiquetas_completas))
    _pintar_encabezado_columnas_ats(ws, fila, etiquetas_completas[:num_columnas_total], color_tema)
    fila += 1
    for i, r in enumerate(filas_normalizadas):
        fila_valores = [r.get("col1"), r.get("col2"), r.get("col3")] + list(r.get("valores", []))
        fila_valores += [None] * (num_columnas_total - len(fila_valores))
        _pintar_fila_datos_ats(ws, fila, fila_valores[:num_columnas_total], es_par=(i % 2 == 0))
        fila += 1
    if total_valores:
        fila_total = [None, "TOTAL", None] + list(total_valores)
        fila_total += [None] * (num_columnas_total - len(fila_total))
        _pintar_fila_total_ats(ws, fila, fila_total[:num_columnas_total], num_columnas_total)
        fila += 1
    fila += 1
    return fila

def escribir_hoja_ats(ws, datos):
    meta = datos["Metadatos"]
    NUM_COLS = 7

    fila = 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=NUM_COLS)
    celda_titulo = ws.cell(row=fila, column=1, value="TALÓN RESUMEN - ANEXO TRANSACCIONAL (ATS)")
    celda_titulo.font = Font(bold=True, size=13, color="FFFFFF")
    celda_titulo.fill = PatternFill(start_color=COLOR_TITULO_ATS, end_color=COLOR_TITULO_ATS, fill_type="solid")
    celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = 24
    fila += 2

    campos_meta = [
        ("Razón Social", meta.get("Razon_Social", "No detectado")),
        ("RUC", meta.get("RUC", "No detectado")),
        ("Período", meta.get("Periodo", "No detectado")),
        ("Fecha de Generación", meta.get("Fecha_Generacion", "No detectado")),
        ("Estado", meta.get("Estado", "No detectado")),
        ("Secuencial Anexo", meta.get("Secuencial_Anexo", "No detectado")),
    ]
    for etiqueta, valor in campos_meta:
        if valor == "No detectado" and etiqueta in ("Estado", "Secuencial Anexo"):
            continue
        celda_etq = ws.cell(row=fila, column=1, value=etiqueta)
        celda_etq.font = Font(bold=True, size=9)
        celda_etq.fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=NUM_COLS)
        celda_val = ws.cell(row=fila, column=2, value=valor)
        celda_val.font = Font(size=9)
        fila += 1
    fila += 1

    for titulo_seccion, clave_seccion in [
        ("COMPRAS", "Compras"),
        ("VENTAS", "Ventas"),
        ("EXPORTACIONES", "Exportaciones"),
    ]:
        info = datos.get(clave_seccion)
        if not info:
            continue
        filas_norm = [
            {"col1": f["codigo"], "col2": f["descripcion"], "col3": f["registros"], "valores": f["valores"]}
            for f in info["filas"]
        ]
        fila = _escribir_seccion_ats(
            ws, fila, titulo_seccion, info["etiquetas_valor"], filas_norm, info["total"],
            COLOR_ENCABEZADO_TRANSACCION, NUM_COLS
        )

    info_renta = datos.get("RetencionRenta")
    if info_renta:
        filas_norm = [
            {"col1": f["codigo"], "col2": f["descripcion"], "col3": f["registros"], "valores": f["valores"]}
            for f in info_renta["filas"]
        ]
        fila = _escribir_seccion_ats(
            ws, fila, "RETENCIÓN EN LA FUENTE DE IMPUESTO A LA RENTA",
            info_renta["etiquetas_valor"], filas_norm, info_renta["total"],
            COLOR_ENCABEZADO_RETENCION_RENTA, NUM_COLS
        )

    info_iva = datos.get("RetencionIVA")
    if info_iva:
        filas_norm = [
            {"col1": f["operacion"], "col2": f["concepto"], "col3": None, "valores": [f["valor"]]}
            for f in info_iva["filas"]
        ]
        fila = _escribir_seccion_ats(
            ws, fila, "RETENCIÓN EN LA FUENTE DE IVA",
            ["Valor Retenido"], filas_norm, info_iva["total"],
            COLOR_ENCABEZADO_RETENCION_IVA, NUM_COLS
        )

    anchos = [12, 46, 12, 16, 16, 16, 16]
    for idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[_columna_letra(idx)].width = ancho
    ws.freeze_panes = "A1"

def generar_nombre_hoja_ats(datos, usados):
    meta = datos["Metadatos"]
    periodo = re.sub(r'[\[\]\:\*\?\/\\]', '', str(meta.get("Periodo", "ATS")))[:14]
    ruc_corto = meta.get("RUC", "")[-4:] if meta.get("RUC", "No detectado") != "No detectado" else ""
    nombre = f"ATS_{periodo}_{ruc_corto}".strip('_')
    nombre = re.sub(r'\s+', '_', nombre)
    if not nombre:
        nombre = "ATS"
    nombre = nombre[:28]
    candidato = nombre
    i = 1
    while candidato in usados:
        i += 1
        candidato = f"{nombre[:25]}_{i}"
    usados.add(candidato)
    return candidato[:31]

def exportar_ats_a_excel():
    if not DATOS_ATS_PROCESADOS:
        safe_messagebox("warning", "Sin datos", "Cargue al menos un PDF de Talón Resumen ATS antes de exportar.")
        return
    usar_existente = messagebox.askyesno(
        "Exportar ATS a Excel",
        "¿Desea agregar los ATS a un archivo Excel YA EXISTENTE?\n\n"
        "Sí = seleccionar un Excel existente y agregarle pestañas nuevas.\n"
        "No = crear un archivo Excel nuevo."
    )
    ruta_existente = None
    if usar_existente:
        ruta_existente = filedialog.askopenfilename(
            title="Seleccionar Excel existente",
            filetypes=[("Libros de Excel", "*.xlsx")]
        )
        if not ruta_existente:
            return
        try:
            wb = openpyxl.load_workbook(ruta_existente)
        except Exception as e:
            safe_messagebox("error", "Error", f"No se pudo abrir el Excel seleccionado: {e}")
            return
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    usados = set(wb.sheetnames)
    hojas_creadas = 0
    errores = []
    for datos in DATOS_ATS_PROCESADOS:
        try:
            nombre_hoja = generar_nombre_hoja_ats(datos, usados)
            ws = wb.create_sheet(title=nombre_hoja)
            escribir_hoja_ats(ws, datos)
            hojas_creadas += 1
        except Exception as e:
            errores.append(f"{os.path.basename(datos['archivo'])}: {e}")

    ruta_guardar = filedialog.asksaveasfilename(
        title="Guardar archivo Excel",
        defaultextension=".xlsx",
        initialfile=os.path.basename(ruta_existente) if ruta_existente else "Talones_ATS.xlsx",
        filetypes=[("Libro de Excel", "*.xlsx")]
    )
    if not ruta_guardar:
        return
    try:
        wb.save(ruta_guardar)
    except Exception as e:
        safe_messagebox("error", "Error", f"No se pudo guardar el archivo: {e}")
        return
    mensaje = f"Se agregaron {hojas_creadas} pestaña(s) de Talón Resumen ATS al archivo:\n{ruta_guardar}"
    if errores:
        mensaje += "\n\nAlgunos archivos fallaron:\n" + "\n".join(errores)
    safe_messagebox("info", "Exportación Exitosa", mensaje)

# ============= INTERFAZ GRÁFICA =============
def obtener_datos_pestana_actual():
    if not DATOS_PROCESADOS:
        return None
    try:
        idx = notebook.index(notebook.select())
    except Exception:
        return None
    if 0 <= idx < len(DATOS_PROCESADOS):
        return DATOS_PROCESADOS[idx]
    return None

def crear_pestana(datos):
    frame = tk.Frame(notebook, bg="white")
    nombre_corto = os.path.basename(datos["archivo"])
    if len(nombre_corto) > 22:
        nombre_corto = nombre_corto[:19] + "..."
    notebook.add(frame, text=nombre_corto)
    val = datos["Validacion"]
    color_estado = "#1b5e20" if val["completo"] else "#c62828"
    texto_estado = "✔ EXTRACCIÓN COMPLETA Y VALIDADA" if val["completo"] else "⚠ REVISAR: FALTAN DATOS CLAVE"
    lbl_estado = tk.Label(frame, text=texto_estado, font=("Tahoma", 10, "bold"), fg=color_estado, bg="white")
    lbl_estado.pack(pady=(8, 4), anchor="w", padx=10)
    if val.get("casilleros_a_revisar", 0) > 0:
        tk.Label(frame, text=f"⚠ {val['casilleros_a_revisar']} casillero(s) leídos por OCR con baja confianza: verificar contra el PDF",
                 font=("Tahoma", 8, "bold"), fg="#e65100", bg="white").pack(anchor="w", padx=10)
    frame_check = tk.Frame(frame, bg="white")
    frame_check.pack(fill="x", padx=10)
    checks = [
        ("RUC", val["ruc_ok"]),
        ("Razón Social", val["razon_ok"]),
        ("Nombre del Formulario", val["formulario_ok"]),
        ("Período Fiscal", val["periodo_ok"]),
        ("Número Serial", val["serial_ok"]),
        ("Casilleros (>0)", val["casilleros_ok"]),
    ]
    for i, (etiqueta, ok) in enumerate(checks):
        simbolo = "✔" if ok else "✘"
        color = "#1b5e20" if ok else "#c62828"
        tk.Label(frame_check, text=f"{simbolo} {etiqueta}", font=("Tahoma", 8, "bold"), fg=color, bg="white")\
            .grid(row=i // 3, column=i % 3, sticky="w", padx=6, pady=2)
    tree_frame = tk.Frame(frame)
    tree_frame.pack(fill="both", expand=True, padx=10, pady=8)
    tree = ttk.Treeview(tree_frame, columns=("Campo", "Valor"), show="headings")
    tree.heading("Campo", text="Componente / Casillero")
    tree.heading("Valor", text="Valor Detectado")
    tree.column("Campo", width=220, anchor="w")
    tree.column("Valor", width=260, anchor="w")
    scroll_y = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=scroll_y.set)
    tree.pack(side="left", fill="both", expand=True)
    scroll_y.pack(side="right", fill="y")
    tree.tag_configure("seccion", background="#e3f2fd", font=("Tahoma", 9, "bold"))
    tree.insert("", "end", values=("── METADATOS ──", ""), tags=("seccion",))
    for campo, contenido in datos["Metadatos"].items():
        tree.insert("", "end", values=(campo, contenido))
    tree.insert("", "end", values=("── CASILLEROS ──", f"Total: {len(datos['Casilleros'])}"), tags=("seccion",))
    for codigo in sorted(datos["Casilleros"].keys(), key=_orden_casillero):
        tree.insert("", "end", values=(f"Casillero {codigo}", datos["Casilleros"][codigo]))
    notebook.select(frame)
    return frame

def refrescar_resumen_global():
    total = len(DATOS_PROCESADOS)
    completos = sum(1 for d in DATOS_PROCESADOS if d["Validacion"]["completo"])
    lbl_resumen.config(
        text=f"Formularios cargados: {total}   |   Completos y validados: {completos}",
        fg="#1b5e20" if completos == total and total > 0 else "#e65100"
    )
    estado_botones = "normal" if total > 0 else "disabled"
    btn_copiar.config(state=estado_botones)
    btn_exportar.config(state=estado_botones)
    btn_quitar.config(state=estado_botones)
    btn_limpiar.config(state=estado_botones)
    btn_reporte_especial.config(state=estado_botones)

def procesar_un_pdf(ruta):
    datos = extraer_datos_formulario_sri(ruta)
    DATOS_PROCESADOS.append(datos)
    crear_pestana(datos)

def procesar_multiples_pdfs(rutas):
    if not rutas:
        return
    lbl_status.config(text=f"PROCESANDO {len(rutas)} ARCHIVO(S)...", fg="#e65100")
    ventana.update_idletasks()
    errores = []
    for ruta in rutas:
        try:
            procesar_un_pdf(ruta)
        except Exception as e:
            errores.append(f"{os.path.basename(ruta)}: {e}")
    refrescar_resumen_global()
    if errores:
        lbl_status.config(text="COMPLETADO CON ADVERTENCIAS", fg="#c62828")
        safe_messagebox("warning", "Algunos archivos fallaron", "\n".join(errores))
    else:
        lbl_status.config(text="DATOS EXTRAÍDOS - LISTOS PARA COPIAR O EXPORTAR", fg="#1b5e20")

def al_soltar_pdf(event):
    if verificar_licencia() is False:
        return
    rutas = limpiar_rutas_arrastre(event.data)
    procesar_multiples_pdfs(rutas)

def examinar_pdfs():
    if verificar_licencia() is False:
        return
    rutas = filedialog.askopenfilenames(title="Seleccionar uno o varios formularios PDF", filetypes=[("Formularios PDF SRI", "*.pdf")])
    procesar_multiples_pdfs(list(rutas))

def quitar_pestana_actual():
    if not DATOS_PROCESADOS:
        return
    try:
        idx = notebook.index(notebook.select())
    except Exception:
        return
    tab_id = notebook.select()
    notebook.forget(tab_id)
    del DATOS_PROCESADOS[idx]
    refrescar_resumen_global()
    lbl_status.config(text="PESTAÑA ELIMINADA" if DATOS_PROCESADOS else "SISTEMA LISTO - ESPERANDO DOCUMENTO", fg="#757575")

def limpiar_todo():
    if not DATOS_PROCESADOS:
        return
    if not messagebox.askyesno("Confirmar", "¿Desea quitar todos los formularios cargados?"):
        return
    for tab_id in notebook.tabs():
        notebook.forget(tab_id)
    DATOS_PROCESADOS.clear()
    refrescar_resumen_global()
    lbl_status.config(text="SISTEMA LISTO - ESPERANDO DOCUMENTO", fg="#757575")

# ============= FUNCIONES ATS =============
def refrescar_resumen_ats():
    total = len(DATOS_ATS_PROCESADOS)
    lbl_resumen_ats.config(text=f"ATS cargados: {total}")
    estado = "normal" if total > 0 else "disabled"
    btn_exportar_ats.config(state=estado)
    btn_quitar_ats.config(state=estado)
    btn_limpiar_ats.config(state=estado)

def examinar_pdfs_ats():
    if verificar_licencia() is False:
        return
    rutas = filedialog.askopenfilenames(
        title="Seleccionar uno o varios Talón Resumen ATS (PDF)",
        filetypes=[("Talón Resumen ATS (PDF)", "*.pdf")]
    )
    if not rutas:
        return
    errores = []
    for ruta in rutas:
        try:
            datos = extraer_datos_ats(ruta)
            DATOS_ATS_PROCESADOS.append(datos)
            listbox_ats.insert("end", os.path.basename(ruta))
        except Exception as e:
            errores.append(f"{os.path.basename(ruta)}: {e}")
    refrescar_resumen_ats()
    if errores:
        safe_messagebox("warning", "Algunos ATS fallaron", "\n".join(errores))

def quitar_ats_seleccionado():
    seleccion = listbox_ats.curselection()
    if not seleccion:
        safe_messagebox("warning", "Sin selección", "Seleccione primero un ATS de la lista.")
        return
    idx = seleccion[0]
    listbox_ats.delete(idx)
    del DATOS_ATS_PROCESADOS[idx]
    refrescar_resumen_ats()

def limpiar_ats_todo():
    if not DATOS_ATS_PROCESADOS:
        return
    if not messagebox.askyesno("Confirmar", "¿Desea quitar todos los ATS cargados?"):
        return
    listbox_ats.delete(0, "end")
    DATOS_ATS_PROCESADOS.clear()
    refrescar_resumen_ats()

# ============= CONSTRUCCIÓN DE LA VENTANA =============
if __name__ == "__main__":
    ventana = TkinterDnD.Tk()
    ventana.title("Sistema de Gestión Tributaria - Extractor de Precisión SRI (v3.0)")
    ventana.geometry("980x720")
    ventana.minsize(900, 650)
    ventana.resizable(True, True)

    tk.Label(ventana, text="Mapeador de Formularios SRI (101, 102, 103, 104, 115 y otros)",
             font=("Tahoma", 13, "bold"), fg="#0d47a1").pack(pady=(12, 4))

    frame_carga = tk.LabelFrame(ventana, text=" Entrada de Documentos ", font=("Tahoma", 9, "bold"), padx=10, pady=8)
    frame_carga.pack(fill="x", padx=15, pady=5)
    tk.Button(frame_carga, text="Examinar PDF(s)", command=examinar_pdfs,
              font=("Tahoma", 10, "bold"), bg="#e3f2fd", relief=tk.GROOVE, cursor="hand2").pack(side="left", padx=5, pady=4)
    tk.Label(frame_carga, text="Puede seleccionar o arrastrar uno o varios PDF a la vez",
             font=("Tahoma", 9, "italic"), fg="gray").pack(side="left", padx=15)

    lbl_status = tk.Label(ventana, text="SISTEMA LISTO - ESPERANDO DOCUMENTO", font=("Tahoma", 10, "bold"), fg="#757575")
    lbl_status.pack(pady=(4, 2))

    lbl_resumen = tk.Label(ventana, text="Formularios cargados: 0   |   Completos y validados: 0",
                            font=("Tahoma", 9, "bold"), fg="#757575")
    lbl_resumen.pack(pady=(0, 6))

    frame_scroll_contenedor = tk.Frame(ventana)
    frame_scroll_contenedor.pack(fill="both", expand=True, padx=15, pady=5)
    canvas_principal = tk.Canvas(frame_scroll_contenedor, highlightthickness=0)
    scrollbar_vertical = ttk.Scrollbar(frame_scroll_contenedor, orient="vertical", command=canvas_principal.yview)
    canvas_principal.configure(yscrollcommand=scrollbar_vertical.set)
    canvas_principal.pack(side="left", fill="both", expand=True)
    scrollbar_vertical.pack(side="right", fill="y")

    frame_interno = tk.Frame(canvas_principal)
    ventana_interna_id = canvas_principal.create_window((0, 0), window=frame_interno, anchor="nw")

    def _actualizar_scrollregion(event=None):
        canvas_principal.configure(scrollregion=canvas_principal.bbox("all"))

    def _ajustar_ancho_interno(event):
        canvas_principal.itemconfig(ventana_interna_id, width=event.width)

    frame_interno.bind("<Configure>", _actualizar_scrollregion)
    canvas_principal.bind("<Configure>", _ajustar_ancho_interno)

    def _rueda_mouse(event):
        canvas_principal.yview_scroll(int(-1 * (event.delta / 120)), "units")

    canvas_principal.bind_all("<MouseWheel>", _rueda_mouse)
    canvas_principal.bind_all("<Button-4>", lambda e: canvas_principal.yview_scroll(-1, "units"))
    canvas_principal.bind_all("<Button-5>", lambda e: canvas_principal.yview_scroll(1, "units"))

    frame_notebook = tk.LabelFrame(frame_interno, text=" Panel de Control y Validación por Formulario ",
                                    font=("Tahoma", 9, "bold"), padx=8, pady=8)
    frame_notebook.pack(fill="both", expand=True, pady=5)

    notebook = ttk.Notebook(frame_notebook)
    notebook.pack(fill="both", expand=True)
    frame_notebook.pack_propagate(True)
    notebook.configure(height=350)

    frame_modo = tk.LabelFrame(frame_interno, text=" Formato de Exportación a Excel ",
                                font=("Tahoma", 9, "bold"), padx=10, pady=8)
    frame_modo.pack(fill="x", pady=(5, 0))

    var_modo_exportacion = tk.StringVar(value="pestanas")
    tk.Radiobutton(
        frame_modo,
        text="Una pestaña por formulario (como hasta ahora)",
        variable=var_modo_exportacion, value="pestanas",
        font=("Tahoma", 9)
    ).pack(anchor="w", padx=6, pady=2)
    tk.Radiobutton(
        frame_modo,
        text="Todos los formularios en la MISMA pestaña, en columnas (1 columna en blanco entre cada uno)",
        variable=var_modo_exportacion, value="columnas",
        font=("Tahoma", 9)
    ).pack(anchor="w", padx=6, pady=2)

    frame_acciones = tk.LabelFrame(frame_interno, text=" Acciones ", font=("Tahoma", 9, "bold"), padx=10, pady=10)
    frame_acciones.pack(fill="x", pady=12)

    btn_copiar = tk.Button(frame_acciones, text="COPIAR PESTAÑA ACTUAL (pegar donde quiera)",
                            command=copiar_al_portapapeles_formateado, font=("Tahoma", 10, "bold"),
                            bg="#fff9c4", fg="#5d4037", relief=tk.RAISED, cursor="hand2", state="disabled")
    btn_copiar.grid(row=0, column=0, sticky="ew", padx=4, pady=4)

    btn_exportar = tk.Button(frame_acciones, text="EXPORTAR TODO A EXCEL (según formato elegido arriba)",
                              command=exportar_a_excel, font=("Tahoma", 10, "bold"),
                              bg="#c8e6c9", fg="#1b5e20", relief=tk.RAISED, cursor="hand2", state="disabled")
    btn_exportar.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

    btn_reporte_especial = tk.Button(frame_acciones, text="REPORTE ESPECIAL (Casilleros × Meses)",
                                      command=generar_reporte_especial_casilleros, font=("Tahoma", 10, "bold"),
                                      bg="#E8B4FF", fg="#4A148C", relief=tk.RAISED, cursor="hand2", state="disabled")
    btn_reporte_especial.grid(row=0, column=2, sticky="ew", padx=4, pady=4)

    btn_quitar = tk.Button(frame_acciones, text="QUITAR PESTAÑA ACTUAL",
                            command=quitar_pestana_actual, font=("Tahoma", 9),
                            bg="#ffe0b2", relief=tk.GROOVE, cursor="hand2", state="disabled")
    btn_quitar.grid(row=1, column=0, sticky="ew", padx=4, pady=4)

    btn_limpiar = tk.Button(frame_acciones, text="LIMPIAR TODO",
                             command=limpiar_todo, font=("Tahoma", 9),
                             bg="#ffcdd2", relief=tk.GROOVE, cursor="hand2", state="disabled")
    btn_limpiar.grid(row=1, column=1, sticky="ew", padx=4, pady=4)

    frame_acciones.columnconfigure(0, weight=1)
    frame_acciones.columnconfigure(1, weight=1)
    frame_acciones.columnconfigure(2, weight=1)

    frame_ats = tk.LabelFrame(frame_interno, text=" Vouchear Talón Resumen ATS (Anexo Transaccional) ",
                               font=("Tahoma", 9, "bold"), padx=10, pady=8)
    frame_ats.pack(fill="x", pady=(0, 12))

    tk.Label(frame_ats, text="Compras/Ventas/Exportaciones y Retenciones de Renta e IVA, con formato a color",
              font=("Tahoma", 8, "italic"), fg="gray").pack(anchor="w", padx=2, pady=(0, 4))

    frame_ats_botones = tk.Frame(frame_ats)
    frame_ats_botones.pack(fill="x")

    tk.Button(frame_ats_botones, text="Examinar PDF(s) ATS", command=examinar_pdfs_ats,
              font=("Tahoma", 9, "bold"), bg="#e1f5fe", relief=tk.GROOVE, cursor="hand2").pack(side="left", padx=4)

    lbl_resumen_ats = tk.Label(frame_ats_botones, text="ATS cargados: 0", font=("Tahoma", 9, "bold"), fg="#757575")
    lbl_resumen_ats.pack(side="left", padx=10)

    frame_ats_lista = tk.Frame(frame_ats)
    frame_ats_lista.pack(fill="x", pady=(6, 6))

    listbox_ats = tk.Listbox(frame_ats_lista, height=4)
    listbox_ats.pack(side="left", fill="x", expand=True)

    scroll_ats = ttk.Scrollbar(frame_ats_lista, orient="vertical", command=listbox_ats.yview)
    listbox_ats.configure(yscrollcommand=scroll_ats.set)
    scroll_ats.pack(side="right", fill="y")

    frame_ats_acciones = tk.Frame(frame_ats)
    frame_ats_acciones.pack(fill="x")

    btn_quitar_ats = tk.Button(frame_ats_acciones, text="QUITAR SELECCIONADO", command=quitar_ats_seleccionado,
                                font=("Tahoma", 9), bg="#ffe0b2", relief=tk.GROOVE, cursor="hand2", state="disabled")
    btn_quitar_ats.pack(side="left", padx=4, pady=4)

    btn_limpiar_ats = tk.Button(frame_ats_acciones, text="LIMPIAR TODO", command=limpiar_ats_todo,
                                 font=("Tahoma", 9), bg="#ffcdd2", relief=tk.GROOVE, cursor="hand2", state="disabled")
    btn_limpiar_ats.pack(side="left", padx=4, pady=4)

    btn_exportar_ats = tk.Button(frame_ats_acciones, text="EXPORTAR ATS A EXCEL (bonito, con colores)",
                                  command=exportar_ats_a_excel, font=("Tahoma", 9, "bold"),
                                  bg="#c8e6c9", fg="#1b5e20", relief=tk.GROOVE, cursor="hand2", state="disabled")
    btn_exportar_ats.pack(side="left", padx=4, pady=4)

    ventana.drop_target_register(DND_FILES)
    ventana.dnd_bind('<<Drop>>', al_soltar_pdf)

    tk.Label(frame_interno, text="Propiedad de Roberto Robles © Unidad Técnica PBP | v3.0",
             font=("Tahoma", 8, "bold"), fg="#b0bec5").pack(anchor="e", padx=10, pady=6)

    ventana.mainloop()
